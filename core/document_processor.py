"""
Document Processing Pipeline
Handles loading, parsing, and chunking of enterprise documents (PDF, DOCX, TXT).
"""

import logging
from pathlib import Path
from typing import Optional

from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Processes enterprise documents into chunks suitable for embedding."""

    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 100):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def load_document(self, file_path: str) -> list[Document]:
        """Load a single document from file path."""
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type: {ext}. "
                f"Supported: {self.SUPPORTED_EXTENSIONS}"
            )

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        logger.info(f"Loading document: {path.name}")

        if ext == ".pdf":
            return self._load_pdf(path)
        elif ext == ".docx":
            return self._load_docx(path)
        elif ext in (".txt", ".md"):
            return self._load_text(path)

        return []

    def _load_pdf(self, path: Path) -> list[Document]:
        """Load a PDF file."""
        from langchain_community.document_loaders import PyPDFLoader

        loader = PyPDFLoader(str(path))
        documents = loader.load()
        for doc in documents:
            doc.metadata["source"] = path.name
            doc.metadata["file_type"] = "pdf"
        return documents

    def _load_docx(self, path: Path) -> list[Document]:
        """Load a DOCX file."""
        from docx import Document as DocxDocument

        docx_doc = DocxDocument(str(path))
        text = "\n".join(
            paragraph.text
            for paragraph in docx_doc.paragraphs
            if paragraph.text.strip()
        )
        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_type": "docx",
                },
            )
        ]

    def _load_text(self, path: Path) -> list[Document]:
        """Load a plain text or markdown file."""
        text = path.read_text(encoding="utf-8")
        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "file_type": path.suffix.lstrip("."),
                },
            )
        ]

    def chunk_documents(self, documents: list[Document]) -> list[Document]:
        """Split documents into smaller chunks for embedding."""
        chunks = self.text_splitter.split_documents(documents)
        logger.info(
            f"Split {len(documents)} document(s) into {len(chunks)} chunks"
        )

        # Add chunk metadata
        for i, chunk in enumerate(chunks):
            chunk.metadata["chunk_index"] = i
            chunk.metadata["chunk_size"] = len(chunk.page_content)

        return chunks

    def process_file(self, file_path: str) -> list[Document]:
        """End-to-end processing: load → chunk a single file."""
        documents = self.load_document(file_path)
        return self.chunk_documents(documents)

    def process_directory(self, directory_path: str) -> list[Document]:
        """Process all supported documents in a directory."""
        dir_path = Path(directory_path)

        if not dir_path.exists():
            raise FileNotFoundError(f"Directory not found: {directory_path}")

        all_chunks: list[Document] = []

        for ext in self.SUPPORTED_EXTENSIONS:
            for file_path in dir_path.glob(f"*{ext}"):
                try:
                    chunks = self.process_file(str(file_path))
                    all_chunks.extend(chunks)
                    logger.info(f"Processed {file_path.name}: {len(chunks)} chunks")
                except Exception as e:
                    logger.error(f"Error processing {file_path.name}: {e}")

        logger.info(
            f"Total: processed {len(all_chunks)} chunks from {directory_path}"
        )
        return all_chunks

    def get_supported_extensions(self) -> set[str]:
        """Return set of supported file extensions."""
        return self.SUPPORTED_EXTENSIONS.copy()
